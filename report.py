"""Live Excel trade report: WON / LOST and percentage per coin."""

from __future__ import annotations

import atexit
import copy
import csv
import logging
import time
from datetime import datetime
from pathlib import Path
from threading import Event, Lock, Thread

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import (
    BUY_SOL_AMOUNT,
    MAX_TRADES,
    BREAKEVEN_ARM_PCT,
    STOP_LOSS_PCT,
    TAKE_PROFIT_TRIGGER_PCT,
    TRAILING_STOP_PCT,
)

REPORT_DIR = Path(__file__).resolve().parent / "reports"
EXCEL_PATH = REPORT_DIR / "pumpfun_trades.xlsx"
BLOCKED_CSV = REPORT_DIR / "blocked.csv"
TRADES_CSV = REPORT_DIR / "trades.csv"
FLUSH_DEBOUNCE_SECONDS = 0.2
_LOG = logging.getLogger("pump-report")
TRADE_CSV_FIELDS = [
    "name",
    "symbol",
    "result",
    "pnl_pct",
    "mfe_pct",
    "mae_pct",
    "invested",
    "exit_sol",
    "pnl_sol",
    "reason",
    "mint",
    "bought_at",
    "sold_at",
]

BLOCK_HEADERS = [
    "Time",
    "Coin",
    "Symbol",
    "LP tokens locked %",
    "Total liquidity USD",
    "Reasons",
    "Top holder tokens",
    "Mint",
]

HEADERS = [
    "#",
    "Coin",
    "Symbol",
    "Result",
    "PnL %",
    "MFE %",
    "MAE %",
    "Invested SOL",
    "Exit SOL",
    "Profit/Loss SOL",
    "Reason",
    "Mint",
    "Bought at",
    "Sold at",
]

FILL_HEADER = PatternFill("solid", fgColor="1F4E79")
FILL_WON = PatternFill("solid", fgColor="C6EFCE")
FILL_LOST = PatternFill("solid", fgColor="FFC7CE")
FILL_OPEN = PatternFill("solid", fgColor="FFF2CC")
FONT_HEADER = Font(bold=True, color="FFFFFF")
FONT_WON = Font(bold=True, color="006100")
FONT_LOST = Font(bold=True, color="9C0006")
FONT_OPEN = Font(bold=True, color="9C5700")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _pct(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:+.1f}%"


_CSV_HEADER_CHECKED = False


def _migrate_trade_csv() -> None:
    """Bring an older trades.csv up to the current columns without dropping rows."""
    global _CSV_HEADER_CHECKED
    if _CSV_HEADER_CHECKED:
        return
    _CSV_HEADER_CHECKED = True
    if not TRADES_CSV.exists():
        return
    try:
        with TRADES_CSV.open("r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle)
            header = reader.fieldnames or []
            rows = list(reader)
    except Exception:
        return
    if header == TRADE_CSV_FIELDS:
        return
    if not set(header) <= set(TRADE_CSV_FIELDS):
        # Unknown columns: keep the old file beside the new one instead of guessing.
        backup = REPORT_DIR / f"trades_{datetime.now():%Y%m%d_%H%M%S}.csv"
        try:
            TRADES_CSV.rename(backup)
            _LOG.info("Unknown trade CSV columns; previous file kept at %s", backup.name)
        except Exception:
            _LOG.exception("Could not rotate the outdated trade CSV")
        return
    try:
        with TRADES_CSV.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=TRADE_CSV_FIELDS)
            writer.writeheader()
            for row in rows:
                writer.writerow({key: row.get(key, "") or "" for key in TRADE_CSV_FIELDS})
        _LOG.info("Migrated trade CSV to %s columns (%s rows kept)", len(TRADE_CSV_FIELDS), len(rows))
    except Exception:
        _LOG.exception("Trade CSV migration failed")


class TradeReport:
    def __init__(self) -> None:
        self._lock = Lock()
        self._rows: list[dict[str, object]] = []
        self._blocked: list[dict[str, object]] = []
        self._dirty = False
        self._stop = Event()
        self._wake = Event()
        REPORT_DIR.mkdir(exist_ok=True)
        self._path = EXCEL_PATH
        self._writer = Thread(target=self._writer_loop, name="excel-writer", daemon=True)
        self._writer.start()
        atexit.register(self.flush)

    @property
    def path(self) -> Path:
        return self._path

    def opened_count(self) -> int:
        return len(self._rows)

    def record_open(
        self,
        *,
        name: str,
        symbol: str,
        mint: str,
        invested_sol: float,
    ) -> None:
        with self._lock:
            self._rows.append(
                {
                    "name": name or symbol,
                    "symbol": symbol,
                    "result": "OPEN",
                    "pnl_pct": None,
                    "mfe_pct": None,
                    "mae_pct": None,
                    "invested": invested_sol,
                    "exit_sol": None,
                    "pnl_sol": None,
                    "reason": "holding",
                    "mint": mint,
                    "bought_at": _now(),
                    "sold_at": "",
                }
            )
            snapshot = dict(self._rows[-1])
            self._mark_dirty()
        self._append_trade_csv(snapshot)

    def record_close(
        self,
        *,
        mint: str,
        pnl_pct: float,
        exit_sol: float,
        reason: str,
        mfe_pct: float | None = None,
        mae_pct: float | None = None,
    ) -> str:
        result = "WON" if pnl_pct >= 0 else "LOST"
        snapshot: dict[str, object] | None = None
        with self._lock:
            for row in reversed(self._rows):
                if str(row["mint"]) == mint and row["result"] == "OPEN":
                    invested = float(row["invested"] or 0)
                    row["result"] = result
                    row["pnl_pct"] = pnl_pct
                    row["mfe_pct"] = mfe_pct
                    row["mae_pct"] = mae_pct
                    row["exit_sol"] = exit_sol
                    row["pnl_sol"] = exit_sol - invested
                    row["reason"] = reason
                    row["sold_at"] = _now()
                    snapshot = dict(row)
                    break
            self._mark_dirty()
        if snapshot:
            self._append_trade_csv(snapshot)
        return result

    def record_add(self, *, mint: str, extra_sol: float) -> None:
        snapshot: dict[str, object] | None = None
        with self._lock:
            for row in reversed(self._rows):
                if str(row["mint"]) == mint and row["result"] == "OPEN":
                    row["invested"] = float(row["invested"] or 0) + extra_sol
                    snapshot = dict(row)
                    break
            self._mark_dirty()
        if snapshot:
            self._append_trade_csv(snapshot)

    def record_block(
        self,
        *,
        name: str,
        symbol: str,
        mint: str,
        reasons: str,
        liquidity_usd: float,
        top_holder_tokens: float,
        lp_locked_pct: float = 0.0,
    ) -> None:
        row = {
            "time": _now(),
            "name": name or symbol,
            "symbol": symbol,
            "reasons": reasons,
            "liquidity_usd": liquidity_usd,
            "lp_locked_pct": lp_locked_pct,
            "top_holder_tokens": top_holder_tokens,
            "mint": mint,
        }
        with self._lock:
            self._blocked.append(row)
            self._blocked = self._blocked[-250:]
            self._mark_dirty()
        self._append_blocked_csv(row)

    def _append_blocked_csv(self, row: dict[str, object]) -> None:
        try:
            new_file = not BLOCKED_CSV.exists()
            with BLOCKED_CSV.open("a", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=[
                    "time", "name", "symbol", "lp_locked_pct", "liquidity_usd",
                    "reasons", "top_holder_tokens", "mint",
                ])
                if new_file:
                    writer.writeheader()
                writer.writerow(row)
        except Exception:
            _LOG.exception("Blocked journal append failed")

    def _append_trade_csv(self, row: dict[str, object]) -> None:
        try:
            _migrate_trade_csv()
            new_file = not TRADES_CSV.exists()
            with TRADES_CSV.open("a", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=TRADE_CSV_FIELDS)
                if new_file:
                    writer.writeheader()
                writer.writerow({key: row.get(key, "") for key in TRADE_CSV_FIELDS})
        except Exception:
            _LOG.exception("Trade journal append failed")

    def _mark_dirty(self) -> None:
        self._dirty = True
        self._wake.set()

    def flush(self) -> None:
        with self._lock:
            rows = copy.deepcopy(self._rows)
            blocked = copy.deepcopy(self._blocked)
            self._dirty = False
        if rows or blocked:
            self._save(rows, blocked)

    def _writer_loop(self) -> None:
        while not self._stop.is_set():
            self._wake.wait(timeout=0.5)
            self._wake.clear()
            if self._stop.is_set():
                break
            time.sleep(FLUSH_DEBOUNCE_SECONDS)
            with self._lock:
                if not self._dirty:
                    continue
                rows = copy.deepcopy(self._rows)
                blocked = copy.deepcopy(self._blocked)
                self._dirty = False
            try:
                self._save(rows, blocked)
            except Exception:
                _LOG.exception("Background Excel save failed")

    def _save(self, rows: list[dict[str, object]], blocked: list[dict[str, object]]) -> None:
        wb = Workbook()
        trades = wb.active
        trades.title = "Trades"
        self._write_trades(trades, rows)
        self._write_blocked(wb.create_sheet("Blocked"), blocked)
        self._write_summary(wb.create_sheet("Summary", 0), rows)
        try:
            wb.save(self._path)
        except PermissionError:
            fallback = REPORT_DIR / "pumpfun_trades_live.xlsx"
            wb.save(fallback)
            self._path = fallback

    def _write_blocked(self, ws, blocked: list[dict[str, object]]) -> None:
        ws.freeze_panes = "A2"
        for col, header in enumerate(BLOCK_HEADERS, 1):
            cell = ws.cell(1, col, header)
            cell.fill = FILL_HEADER
            cell.font = FONT_HEADER
            cell.border = THIN
        for idx, row in enumerate(blocked, 1):
            values = [
                row["time"],
                row["name"],
                row["symbol"],
                row.get("lp_locked_pct", 0),
                row["liquidity_usd"],
                row["reasons"],
                row["top_holder_tokens"],
                row["mint"],
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(idx + 1, col, value)
                cell.border = THIN
                cell.fill = FILL_LOST
                if col == 4:
                    cell.number_format = "0.00"
                if col in {5, 7} and value != "":
                    cell.number_format = "#,##0.00"
        widths = [20, 28, 14, 20, 24, 70, 20, 48]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

    def _write_trades(self, ws, rows: list[dict[str, object]]) -> None:
        ws.freeze_panes = "A2"
        last_col = get_column_letter(len(HEADERS))
        ws.auto_filter.ref = f"A1:{last_col}{max(len(rows) + 1, 2)}"
        for col, header in enumerate(HEADERS, 1):
            cell = ws.cell(1, col, header)
            cell.fill = FILL_HEADER
            cell.font = FONT_HEADER
            cell.alignment = Alignment(horizontal="center")
            cell.border = THIN

        for idx, row in enumerate(rows, 1):
            result = str(row["result"])
            fill = FILL_OPEN
            font = FONT_OPEN
            if result == "WON":
                fill, font = FILL_WON, FONT_WON
            elif result == "LOST":
                fill, font = FILL_LOST, FONT_LOST

            pnl_sol = row["pnl_sol"]
            values = [
                idx,
                row["name"],
                row["symbol"],
                result,
                _pct(row["pnl_pct"] if isinstance(row["pnl_pct"], float) else None),
                _pct(row["mfe_pct"] if isinstance(row["mfe_pct"], float) else None),
                _pct(row["mae_pct"] if isinstance(row["mae_pct"], float) else None),
                float(row["invested"] or 0),
                "" if row["exit_sol"] is None else float(row["exit_sol"]),
                "" if pnl_sol is None else float(pnl_sol),
                row["reason"],
                row["mint"],
                row["bought_at"],
                row["sold_at"],
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(idx + 1, col, value)
                cell.border = THIN
                cell.alignment = Alignment(vertical="center")
                if col in {1, 4, 5, 6, 7}:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                if col in {4, 5}:
                    cell.fill = fill
                    cell.font = font
                if col in {8, 9, 10} and value != "":
                    cell.number_format = "0.0000"

        widths = [6, 28, 14, 12, 12, 10, 10, 14, 12, 16, 36, 48, 20, 20]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

    def _write_summary(self, ws, rows: list[dict[str, object]]) -> None:
        closed = [r for r in rows if r["result"] in {"WON", "LOST"}]
        wins = [r for r in closed if r["result"] == "WON"]
        losses = [r for r in closed if r["result"] == "LOST"]
        invested = sum(float(r["invested"] or 0) for r in closed)
        returned = sum(float(r["exit_sol"] or 0) for r in closed)
        pnl = returned - invested
        win_rate = (len(wins) / len(closed) * 100) if closed else 0.0
        mfes = [float(r["mfe_pct"]) for r in closed if isinstance(r.get("mfe_pct"), float)]
        maes = [float(r["mae_pct"]) for r in closed if isinstance(r.get("mae_pct"), float)]
        avg_mfe = sum(mfes) / len(mfes) if mfes else 0.0
        avg_mae = sum(maes) / len(maes) if maes else 0.0

        ws["A1"] = "Pump.fun paper-trade report"
        ws["A1"].font = Font(bold=True, size=16, color="1F4E79")
        ws.merge_cells("A1:B1")

        lines = [
            ("Mode", "DRY-RUN"),
            ("Investment per trade (SOL)", BUY_SOL_AMOUNT),
            ("Stop-loss", f"{STOP_LOSS_PCT:.1f}%"),
            ("Breakeven after", f"+{BREAKEVEN_ARM_PCT:.1f}%"),
            ("Trailing TP arm", f"{TAKE_PROFIT_TRIGGER_PCT:.0f}%"),
            ("Trailing callback", f"{TRAILING_STOP_PCT:.0f}%"),
            ("Target trades", MAX_TRADES),
            ("Opened so far", len(rows)),
            ("Still open", sum(1 for r in rows if r["result"] == "OPEN")),
            ("Closed trades", len(closed)),
            ("WON", len(wins)),
            ("LOST", len(losses)),
            ("Win rate", f"{win_rate:.1f}%"),
            ("Avg peak (MFE)", f"{avg_mfe:+.1f}%"),
            ("Avg drawdown (MAE)", f"{avg_mae:+.1f}%"),
            ("Total invested (SOL)", invested),
            ("Total back (SOL)", returned),
            ("Net P&L (SOL)", pnl),
            ("Excel file", str(EXCEL_PATH)),
        ]
        ws["A3"] = "Item"
        ws["B3"] = "Value"
        ws["A3"].font = FONT_HEADER
        ws["B3"].font = FONT_HEADER
        ws["A3"].fill = FILL_HEADER
        ws["B3"].fill = FILL_HEADER
        for i, (label, value) in enumerate(lines, 4):
            ws.cell(i, 1, label).border = THIN
            cell = ws.cell(i, 2, value)
            cell.border = THIN
            if label == "WON":
                cell.fill = FILL_WON
                cell.font = FONT_WON
            elif label == "LOST":
                cell.fill = FILL_LOST
                cell.font = FONT_LOST
            elif label == "Net P&L (SOL)":
                cell.fill = FILL_WON if float(pnl) >= 0 else FILL_LOST
                cell.number_format = "0.0000"
            elif isinstance(value, float) and "rate" not in label.lower():
                cell.number_format = "0.0000"
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 70


REPORT = TradeReport()
